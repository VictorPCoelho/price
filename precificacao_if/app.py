"""
app.py — Precificação de Crédito IF (v12)
Layout v9: 5 abas de configuração no topo + abas de resultado embaixo
Lógica v12: Modo 1 (taxa mínima como output) + Modo 2 (spread travado)
             Aba Planilha estilo Excel com nomenclatura DIFIN
"""
import streamlit as st
import pandas as pd
import sys, os, io, copy, math

sys.path.insert(0, os.path.dirname(__file__))
from recipe          import RECEITAS, carregar
from engine_core     import executar
from exporters       import para_excel
from curve_parser    import parse_difin, TIPOS_CURVA
from rating_model    import (
    build_tabela, CriteriosViabilidade,
    taxa_de_pct_cdi, pct_cdi_de_taxa,
    RATINGS_ORDEM, CORES_RATING, BACKCORES_RATING,
    PD_DEFAULTS_AA, LGD_DEFAULT,
)
from portfolio_sim   import simular_portfolio, simular_portfolio_v2, _clone_direto
from pe_loader       import carregar_pe_arquivo, gerar_csv_template
from funding_faixas  import (FAIXAS_PADRAO, wacc_faixas, resumo_faixas,
                              taxa_am_de_aa, taxa_total_faixa)
from funding_rating  import (FUNDING_DEFAULTS_AA, funding_am_por_rating,
                              resumo_por_letra, gerar_template_funding_csv as gerar_csv_funding,
                              gerar_template_funding_xlsx as gerar_xlsx_funding,
                              carregar_funding_arquivo)

st.set_page_config(page_title="Precificação IF", page_icon="📊",
                   layout="wide", initial_sidebar_state="collapsed")

st.markdown("""<style>
.block-container{padding-top:.6rem;padding-bottom:1rem;max-width:1600px}
.mod-hdr{color:#fff;border-radius:5px 5px 0 0;padding:5px 10px;font-size:10px;font-weight:700;letter-spacing:.06em}
.mod-body{border:1px solid var(--color-border-tertiary);border-top:none;border-radius:0 0 5px 5px;
          padding:7px;margin-bottom:10px;background:var(--color-background-secondary)}
.slot-lbl{font-size:9px;font-weight:700;color:#888;text-transform:uppercase;letter-spacing:.06em;margin:6px 0 3px}
.fcrd{border:1.5px solid var(--color-border-tertiary);border-radius:5px;padding:6px 8px;
      margin-bottom:3px;background:var(--color-background-primary);min-height:52px}
.fcrd .cod{font-family:monospace;font-size:10px;font-weight:700}
.fcrd .lbl{font-size:10px;font-weight:600;color:var(--color-text-primary)}
.fcrd .desc{font-size:9px;color:var(--color-text-secondary);line-height:1.3}
.sec-hdr{font-size:10px;font-weight:700;color:#1F3B6E;text-transform:uppercase;
         letter-spacing:.06em;border-bottom:1px solid var(--color-border-tertiary);
         padding-bottom:2px;margin:8px 0 5px}
.badge{display:inline-block;font-size:9px;padding:1px 5px;border-radius:7px;
       font-weight:600;font-family:monospace;background:#EEF2FA;color:#1F3B6E;margin:1px}
.pill-du{background:#EEF2FA;color:#0C447C;font-size:9px;padding:1px 5px;border-radius:7px;font-weight:600}
.pill-dc{background:#FFF3E0;color:#633806;font-size:9px;padding:1px 5px;border-radius:7px;font-weight:600}
.pill-men{background:#EAF3DE;color:#27500A;font-size:9px;padding:1px 5px;border-radius:7px;font-weight:600}
.pill-fix{background:#F1EFE8;color:#444441;font-size:9px;padding:1px 5px;border-radius:7px;font-weight:600}
.front-banner{border-radius:10px;padding:12px 16px;display:flex;align-items:center;gap:14px;margin-bottom:10px}
.decomp-bar{display:flex;height:8px;border-radius:4px;overflow:hidden;width:100%}
.seg-f{background:#378ADD}.seg-p{background:#E24B4A}
.seg-t{background:#EF9F27}.seg-m{background:#3B6D11}
div[data-testid="metric-container"]{background:var(--color-background-secondary);
    border:1px solid var(--color-border-tertiary);border-radius:7px;padding:8px 12px}
.stTabs [data-baseweb="tab"]{font-size:12px;padding:6px 14px}
</style>""", unsafe_allow_html=True)

# ── Mapeamento colunas → nomenclatura DIFIN ───────────────────────────────────
COLUNAS_MANUAL = {
    "t":        ("t",      "Período"),
    "LCA":      ("LCA",    "Liberação de crédito do ativo"),
    "JA":       ("JA",     "Juros do ativo — z036/z15/z16/z17"),
    "EJA":      ("EJA",    "Exig. juros do ativo — z043"),
    "SJA":      ("SJA",    "Saldo juros exig. ativo"),
    "EBA":      ("EBA",    "Encargo básico do ativo — z1/z2/z3"),
    "ECA":      ("ECA",    "Encargo de capital do ativo — z061"),
    "SCA":      ("SCA",    "Saldo de capital do ativo — z063"),
    "PMTA":     ("PMTA",   "Prestação do mutuário ativo — z079"),
    "PE":       ("PE",     "Parcela de equalização"),
    "SDA":      ("SDA",    "Saldo devedor do ativo — z082"),
    "LCP":      ("LCP",    "Liberação de crédito do passivo"),
    "EBP":      ("EBP",    "Encargo básico do passivo — z103"),
    "EEBP":     ("EEBP",   "Exig. encargo básico passivo — z109"),
    "SEBP":     ("SEBP",   "Saldo encargo básico exig. — z113"),
    "ECP":      ("ECP",    "Encargo de capital do passivo — z154"),
    "SCP":      ("SCP",    "Saldo de capital do passivo"),
    "PMTP":     ("PMTP",   "Prestação do mutuário passivo — z161"),
    "SDP":      ("SDP",    "Saldo devedor do passivo — z173"),
    "SP":       ("SP",     "Spread — z175/z176/z177"),
    "tarifa":   ("TAR",    "Tarifa por período"),
    "BFA":      ("BFA",    "Base de faturamento do ativo — z236"),
    "pasep_cof":("PC",     "PASEP + COFINS — z306"),
    "iss":      ("ISS",    "ISS sobre tarifas"),
    "DPE":      ("DPE",    "Despesa de provisão extra — z276/ipp"),
    "IPP1":     ("IPP1",   "Imputação PE Fase 1 — z276"),
    "EPE":      ("EPE",    "Exig. de perda esperada — z444"),
    "SIPP":     ("SIPP",   "Saldo acum. provisão IPP"),
    "MG":       ("MG",     "Margem de ganho — z182/z180/z393"),
    "MG_fla":   ("MG_t-1", "Margem de ganho período anterior"),
    "MC":       ("MC",     "Margem de contribuição — z196"),
    "z394":     ("z394",   "Custo de contratação — z394"),
    "z230":     ("z230",   "Custo de manutenção — z230"),
    "f1046":    ("f1046",  "Custo de agenciamento — f1046"),
    "z236":     ("z236",   "Custo proporcional — z236"),
    "f1044":    ("f1044",  "Custos operacionais totais — f1044"),
    "BFB":      ("BFB",    "Base de faturamento do banco — f1048"),
    "ir_cs":    ("IR/CS",  "IRPJ + CSLL — z355"),
    "FLA":      ("FLA",    "Fluxo A — resultado líquido — f1060"),
    "FLA_VP":   ("FLA_VP", "FLA atualizado a VP — f1062"),
    "FLC":      ("FLC",    "Fluxo C — amortização do ativo"),
    "FLE":      ("FLE",    "Fluxo E — amortização do passivo"),
    "EPE_VP":   ("EPE_VP", "EPE atualizada a VP"),
}

GRUPOS_PLANILHA = {
    "Ativo (Módulo 3)":          ["t","LCA","JA","EJA","SJA","EBA","ECA","SCA","PMTA","PE","SDA"],
    "Passivo (Módulo 4)":        ["t","LCP","EBP","EEBP","SEBP","ECP","SCP","PMTP","SDP"],
    "Spread e Margem (Módulo 5)":["t","SP","MG","MG_fla","MC","BFA"],
    "Provisão e Risco":          ["t","DPE","IPP1","EPE","EPE_VP","SIPP"],
    "Tributos":                  ["t","pasep_cof","iss","ir_cs","BFB"],
    "Custos Operacionais":       ["t","z394","z230","f1046","z236","f1044"],
    "Fluxos (Módulo 7)":         ["t","FLA","FLA_VP","FLC","FLE"],
    "Completo":                  list(COLUNAS_MANUAL.keys()),
}

# ── Catálogo completo de fórmulas ─────────────────────────────────────────────
SLOTS = {
    "mod2":{"cor":"#854F0B","bg":"#FFF3E0","label":"MÓDULO 2 — ENCARGOS","slots":{
        "Taxa ativo (TX^k)":{"chave":"enc.taxa_ativo","formulas":{
            "F1070":("Padrão ★","i^k = TX^k"),
            "F1071":("Taxa efetiva","[(1+TN/n×100)^(ND/NC)−1]×100","men"),
            "F1072":("Taxas equiv.","[(1+TC/100)^(ND/NC)−1]×100","men"),
            "F1073":("Proporcional","i^k=TC_b/(b/g)","fix"),
            "F1076":("Composta","ic^k=ind+i^q+i^k","fix"),
        }},
        "Indexador ativo":{"chave":"enc.ind_ativo","formulas":{
            "F1082":("Padrão ★","ind^k = INDEX^k"),
            "F1083":("A termo","forward DU/252 anual.","du"),
            "F1090":("A termo/mês","forward DU/252 per.","du"),
            "F1197":("Intl. DC/360","a termo intl. 1","dc"),
            "F1198":("Intl. LIBOR","ponder. LIBOR DC","dc"),
        }},
        "Taxa passivo (TX^q)":{"chave":"enc.taxa_passivo","formulas":{
            "F1077":("Padrão ★","i^q = TX^q"),
            "F1079":("Taxas equiv.","[(1+TC/100)^(ND/NC)−1]×100","men"),
            "F1081":("Composta","ic^q=ind^q+i^q","fix"),
            "F1078":("Efetiva","[(1+TN^q/n×100)^(ND/NC)−1]×100","men"),
            "F1102":("Poupança","70%×CVCC se ≤8,5%","men"),
        }},
        "Indexador passivo":{"chave":"enc.ind_passivo","formulas":{
            "F1087":("Padrão ★","ind^q = INDEX^q"),
            "F1088":("A termo","forward DU/252 anual.","du"),
            "F1091":("A termo/mês","forward DU/252 per.","du"),
            "F1195":("Intl. DC/360","a termo intl. 1 passivo","dc"),
            "F1196":("Intl. LIBOR","ponder. LIBOR passivo","dc"),
        }},
    }},
    "ativo":{"cor":"#1F3B6E","bg":"#EEF2FA","label":"MÓDULO 3 — ATIVO","slots":{
        "JA — Juros":{"chave":"ativo.JA","formulas":{
            "z036":("Flat mensal ★","SDA × i"),
            "z15": ("Equiv. 1/n","(1+i)^(1/n)−1","men"),
            "z16": ("DU/252","(1+i)^(DU/252)−1","du"),
            "z17": ("DC/360","(1+i)^(DC/360)−1","dc"),
        }},
        "PMTA":{"chave":"ativo.PMTA","formulas":{
            "price":("Price ★","Prestação fixa"),
            "sac":  ("SAC","Amort. constante"),
            "bullet":("Bullet","Paga em T"),
            "z31":  ("Periódica","a cada Z meses","men"),
        }},
        "SDA":{"chave":"ativo.SDA","formulas":{
            "z45":("Convencional ★","SDA+EBA+JA−PMTA"),
            "z46":("Sem JA sep.","SDA+EBA−PMTA"),
        }},
    }},
    "passivo":{"cor":"#1A5C3A","bg":"#E8F5ED","label":"MÓDULO 4 — PASSIVO","slots":{
        "EBP — Encargo":{"chave":"passivo.EBP","formulas":{
            "ebp_z2":("Equiv. 1/n","(1+i)^(1/n)−1","men"),
            "ebp_z3":("DU/252 ★","(1+i)^(DU/252)−1","du"),
            "ebp_z4":("DC/360","(1+i)^(DC/360)−1","dc"),
            "ebp_z1":("Flat","EBP flat sobre SDP","fix"),
        }},
        "ECP — Capital":{"chave":"passivo.ECP","formulas":{
            "z407":   ("Matched ★","ECP=PMTA−EJA"),
            "ecp_z68":("Bullet","ECP total em T"),
            "ecp_z65":("SAC passivo","SCP_{t-1}/n_rest.","men"),
        }},
    }},
    "fluxo":{"cor":"#4A1080","bg":"#F3EFFE","label":"MÓDULO 5 — FLUXO","slots":{
        "SP — Spread":{"chave":"fluxo.SP","formulas":{
            "z175":("Simples ★","PMTA−PMTP"),
            "z176":("Com equaliz.","PMTA−PMTP+EEQL"),
            "z177":("Multi-ativo","ΣPMTA−ΣPMTP"),
        }},
        "MG — Margem":{"chave":"fluxo.MG","formulas":{
            "z182":("Mensal 1/n ★","(SDA+JA)×((1+λ)^(1/n)−1)","men"),
            "z180":("DU/252","(SDA+JA)×((1+λ)^(DU/252)−1)","du"),
            "z393":("Base Price","(SDA_{t-1}+JA)×λ_du","du"),
            "z181":("DC/360","(SDA+JA)×λ_dc","dc"),
            "z441":("Fixo 1/12","base×((1+λ)^(1/12)−1)","fix"),
        }},
        "DPE — Provisão":{"chave":"fluxo.DPE","formulas":{
            "ipp": ("IPP 2 fases ★","Fase1+Fase2 flat"),
            "z276":("Mensal RBA","base×RBA×r×(1/n)","men"),
            "z277":("DU/252 RBA","base×RBA×r×DU","du"),
            "z278":("DC/365 RBA","base×RBA×r×DC","dc"),
        }},
    }},
    "perf":{"cor":"#3C3489","bg":"#EEEDFE","label":"MÓDULO 6 — PERFORMANCE","slots":{
        "RSPLE":{"chave":"perf.RSPLE","formulas":{
            "F1064":("Padrão ★","λ_aa/(FPR×K×(1+(F−1)×FCC))"),
            "F1223":("Com equaliz.","÷FP","fix"),
        }},
        "RAR":{"chave":"perf.RAR","formulas":{
            "F1065":("Base","λ_mg/CE"),
            "F1183":("Gestão ★","λ_mg/CEIRB — usa Kp"),
            "F1213":("Prudencial","λ_mg/Kp_pond","fix"),
            "F1229":("Sem RA","φ_mga/Kp","fix"),
        }},
        "Iterativas":{"chave":"perf.iter","formulas":{
            "F1095":("TIR ★","Fluxo C"),
            "F1096":("TPP ★","Fluxo E"),
            "F1233":("SPTX ★","Spread em taxa"),
            "F1097":("% CDI DU","Fluxo C DU/252","du"),
            "F1098":("% CDI All In","Fluxo D DU/252","du"),
        }},
    }},
}

DEFAULTS = {
    "enc.taxa_ativo":"F1070","enc.ind_ativo":"F1082",
    "enc.taxa_passivo":"F1077","enc.ind_passivo":"F1087",
    "ativo.JA":"z036","ativo.PMTA":"price","ativo.SDA":"z45",
    "passivo.EBP":"ebp_z3","passivo.ECP":"z407",
    "fluxo.SP":"z175","fluxo.MG":"z182","fluxo.DPE":"ipp",
    "perf.RSPLE":"F1064","perf.RAR":"F1183","perf.iter":"F1095",
}
PILL_MAP={"du":"<span class='pill-du'>DU/252</span>","dc":"<span class='pill-dc'>DC/360</span>",
          "men":"<span class='pill-men'>mensal</span>","fix":"<span class='pill-fix'>flat</span>"}
EXIG_OPTS={"Mensal ★":1,"Bimestral":2,"Trimestral":3,"Semestral":6,"Anual":12,"Bullet":None,"Proporcional":-1}

for k,v in [
    ("slots_sel",dict(DEFAULTS)),("resultado_m1",None),("cenarios",{}),
    ("curvas",None),("pd_upload",None),
    ("faixas_funding",copy.deepcopy(FAIXAS_PADRAO)),("modo_app","modo1"),
    ("funding_por_rating",dict(FUNDING_DEFAULTS_AA)),
    ("fund_uniforme",False),
]:
    if k not in st.session_state: st.session_state[k]=v

# ── Helpers ───────────────────────────────────────────────────────────────────
def _gerar_template_excel():
    import io, openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    tab=build_tabela(); wb=openpyxl.Workbook(); ws=wb.active; ws.title='PE por Rating'
    CORES={'A':'E8F5ED','B':'EEF2FA','C':'F3EFFE','D':'FFF3E0','E':'FCEBEB','F':'FFD9D9','G':'FFCCCC'}
    for i,h in enumerate(['Rating','Letra','PD % a.a.','LGD %','PE % a.a.','PE % a.m.','Obs.'],1):
        c=ws.cell(row=1,column=i,value=h); c.font=Font(bold=True,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor='1F3B6E'); c.alignment=Alignment(horizontal='center')
    for i,cod in enumerate(RATINGS_ORDEM,2):
        cls=tab[cod]
        for j,val in enumerate([cod,cls.letra,round(cls.pd_aa,2),round(cls.lgd,1),round(cls.pe_aa,4),round(cls.pe_am,6),'Edite PD'],1):
            ws.cell(row=i,column=j,value=val).fill=PatternFill('solid',fgColor=CORES.get(cls.letra,'F5F5F5'))
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=13
    buf=io.BytesIO(); wb.save(buf); buf.seek(0); return buf.read()

def _calc_rsple(lamb,FPR,K,F,FCC):
    la=((1+lamb/100)**12-1); d=FPR/100*K/100*(1+(F-1)*FCC/100)
    return la/d*100 if d else 0

def _calc_rar(MGVP,CE,CEIRB,modo="F1183"):
    if modo=="F1065": return MGVP/CE*100 if CE else 0
    return MGVP/CEIRB*100 if CEIRB else 0

def taxa_produtoria_am(custo_aa,spread_aa):
    return taxa_am_de_aa(((1+custo_aa/100)*(1+spread_aa/100)-1)*100)

def taxa_soma_am(custo_aa,spread_aa):
    return taxa_am_de_aa(custo_aa+spread_aa)

def calcular_taxa_min(rc_base, cls_rating, rar_min, rsple_min, cdi_am, funding_am_rating=None):
    """Busca binária: menor taxa que atinge os critérios dado o PE e custo de funding do rating."""
    rc=copy.deepcopy(rc_base); rc['params_risco']['fase1_am']=cls_rating.pe_am
    if funding_am_rating is not None:
        rc['params_op']['cdi_am']=funding_am_rating  # custo específico deste rating
    lo,hi=0.1,6.0
    for _ in range(70):
        mid=(lo+hi)/2; rc['params_op']['i_am']=mid; rc['params_op']['PMTA_fixo']=None
        try:
            rs=executar(rc)['resumo']
            ok=((rs.get('RAR_G % (gestão, F1183)',0) or 0)>=rar_min and
                (rs.get('RSPLE %',0) or 0)>=rsple_min and
                (rs.get('FLA VP (Σ) → 0',-1) or -1)>=0)
        except: ok=False
        if ok: hi=mid
        else: lo=mid
        if hi-lo<0.0005: break
    tx=round(hi,4)
    if tx>=5.9: return None
    rc['params_op']['i_am']=tx; rc['params_op']['PMTA_fixo']=None
    try:
        res=executar(rc); rs=res['resumo']; df=res['fluxo']
        rows=df[df.t>0]; tja=rows['JA'].sum() or 1
        return {
            'rating':cls_rating.codigo,'letra':cls_rating.letra,
            'pd_aa':cls_rating.pd_aa,'pe_am':round(cls_rating.pe_am,6),
            'taxa_min_am':tx,'taxa_min_aa':round(((1+tx/100)**12-1)*100,4),
            'pct_cdi':round(tx/cdi_am*100,1) if cdi_am else 0,
            'RSPLE':round(rs.get('RSPLE %',0),4),
            'RAR_G':round(rs.get('RAR_G % (gestão, F1183)',0),4),
            'TIR':rs.get('TIR % a.a.'),
            'FLA_VP':round(rs.get('FLA VP (Σ) → 0',0),2),
            'Duration':round(rs.get('Duration (meses)',0),2),
            'Total_SP':round(rs.get('Total SP',0),2),
            'Total_MG':round(rs.get('Total MG',0),2),
            'decomp':{'f':round(rows['EBP'].sum()/tja*100,1),'p':round(rows['IPP1'].abs().sum()/tja*100,1),
                      't':round((rows['pasep_cof'].abs().sum()+rows['iss'].abs().sum()+rows['ir_cs'].abs().sum())/tja*100,1),
                      'm':round(rows['MG'].sum()/tja*100,1)},
            '_rs':rs,'_df':df,
        }
    except: return None

def render_slot(sn, slot, cor, bg):
    chave=slot["chave"]; sel=st.session_state.slots_sel.get(chave,list(slot["formulas"].keys())[0])
    st.markdown(f"<div class='slot-lbl'>{sn}</div>",unsafe_allow_html=True)
    ncols=min(len(slot["formulas"]),3); cols_sl=st.columns(ncols)
    for i_s,(cod,fdata) in enumerate(slot["formulas"].items()):
        lbl=fdata[0]; desc=fdata[1]; base=fdata[2] if len(fdata)>2 else ""; ativo=(sel==cod); pill=PILL_MAP.get(base,"")
        with cols_sl[i_s%ncols]:
            st.markdown(
                f"<div class='fcrd' style='border-color:{''+cor if ativo else 'var(--color-border-tertiary)'};"
                f"background:{''+bg if ativo else 'var(--color-background-primary)'}'>"
                f"<div style='display:flex;align-items:center;gap:4px'>"
                f"<span class='cod' style='color:{cor};font-weight:{'700' if ativo else '400'}'>"
                f"{'✓ ' if ativo else ''}{cod}</span>{pill}</div>"
                f"<div class='lbl'>{lbl}</div><div class='desc'>{desc}</div></div>",unsafe_allow_html=True)
            if st.button(f"{'✓' if ativo else '→'}",key=f"b_{chave}_{cod}",
                         use_container_width=True,type="primary" if ativo else "secondary"):
                st.session_state.slots_sel[chave]=cod; st.rerun()

def render_planilha(df, tc_val, label=""):
    if df is None: st.info("Sem fluxo."); return
    df=df.copy()
    df["EPE_VP"]=df.apply(lambda r:abs(r["EPE"])/(1+0.13)**(r["t"]*21/252) if r["t"]>0 else 0,axis=1)
    grupo=st.selectbox("Grupo de colunas",list(GRUPOS_PLANILHA.keys()),key=f"grp_{label}",label_visibility="collapsed")
    cols=[c for c in GRUPOS_PLANILHA[grupo] if c in df.columns]
    df_s=df[cols].copy()
    # Renomeia para código DIFIN
    df_s=df_s.rename(columns={c:COLUNAS_MANUAL[c][0] for c in cols if c in COLUNAS_MANUAL})
    for col in df_s.columns:
        if col=="t": continue
        df_s[col]=df_s[col].apply(lambda v:f"{v:,.4f}" if isinstance(v,float) and abs(v)>1e-9 else ("—" if isinstance(v,float) else str(v)))
    def hl_plan(row):
        t_=df.iloc[row.name]["t"] if row.name<len(df) else 0
        if t_==0: return ["background:#EEF2FA"]*len(row)
        if tc_val>0 and 0<t_<=tc_val: return ["background:#FFF9E6"]*len(row)
        return [""]*len(row)
    st.dataframe(df_s.style.apply(hl_plan,axis=1),hide_index=True,use_container_width=True,height=480)
    with st.expander("Legenda das colunas"):
        rows_leg=[{"Código":COLUNAS_MANUAL[c][0],"Campo":c,"Descrição":COLUNAS_MANUAL[c][1]}
                  for c in cols if c in COLUNAS_MANUAL]
        st.dataframe(pd.DataFrame(rows_leg),hide_index=True,use_container_width=True)
    if tc_val>0: st.caption("🟡 Períodos de carência em amarelo | 🔵 t=0 em azul")

# ── Receita base comum ────────────────────────────────────────────────────────
def _build_rc(produto, C, T, tc, i_am, cdi_ef, lamb, pagar_car, z_J,
              custo_cont, custo_manut, alfa_pc, alfa_iss, alfa_ir,
              FPR, K, Kp, F_aval, FCC, fase1_am, usa_f2, fase2_pct,
              tem_pe, pe_pct, pe_parc):
    rc=copy.deepcopy(carregar(produto))
    rc["params_op"].update(dict(C=C,T=T,tc=tc,i_am=i_am,cdi_am=cdi_ef,PMTA_fixo=None,
        pe_parcela=pe_parc if tem_pe else None,pe_pct=pe_pct/100 if tem_pe else 0.0,
        pagar_juros_carencia=pagar_car,z_J_ativo=z_J,CVSC_aa=13.0,DU=21))
    rc["params_margem"]["lamb_am"]=lamb
    rc["params_custo"].update(dict(custo_cont=custo_cont,custo_manut=custo_manut,
        alfa_pc=alfa_pc,alfa_iss=alfa_iss,alfa_ir_cs=alfa_ir))
    rc["params_risco"].update(dict(FPR=FPR,K=K,Kp=Kp,F=F_aval,FCC=FCC,
        fase1_am=fase1_am,usa_fase2=usa_f2,fase2_pct=fase2_pct))
    for ch,cod in st.session_state.slots_sel.items():
        if "perf" not in ch and "enc" not in ch:
            rc["slots"][ch]=cod
    return rc

# ═══════════════════════════════════════════════════════════════════
# CABEÇALHO
# ═══════════════════════════════════════════════════════════════════
m1a=(st.session_state.modo_app=="modo1")

ha,hb=st.columns([2,1])
with ha: st.markdown("## 📊 Precificação de Crédito — IF")
with hb:
    produto_sel=st.selectbox("Produto",list(RECEITAS.keys()),
        format_func=lambda x:RECEITAS[x]["descricao"][:40],label_visibility="collapsed")

# Seletor de modo — faixa larga, bem visível
st.markdown("**Modo de análise:**")
hm1,hm2,hm3=st.columns([2,2,3])
with hm1:
    if st.button(
        "🧮  Modo 1 — Precificação Individual" + (" ✓" if m1a else ""),
        type="primary" if m1a else "secondary",
        use_container_width=True,
        help="Inputs: custo do funding + λ + PE por rating → Output: taxa mínima por rating"
    ):
        st.session_state.modo_app="modo1"; st.rerun()
with hm2:
    if st.button(
        "🔒  Modo 2 — Spread Travado" + (" ✓" if not m1a else ""),
        type="primary" if not m1a else "secondary",
        use_container_width=True,
        help="Spread IF fixo: custo + spread → taxa ao cliente por faixa × rating"
    ):
        st.session_state.modo_app="modo2"; st.rerun()
with hm3:
    if m1a:
        st.caption("**Modo 1:** você informa os custos e a margem alvo → o sistema calcula a taxa mínima que cada rating precisa pagar.")
    else:
        st.caption("**Modo 2:** o spread da IF é travado → o sistema calcula a taxa ao cliente por faixa e mostra até qual rating é viável.")

receita_base=carregar(produto_sel)
op=receita_base["params_op"]; cst=receita_base["params_custo"]; rsc=receita_base["params_risco"]
st.markdown("---")

# ═══════════════════════════════════════════════════════════════════
# TOPO — 5 ABAS DE CONFIGURAÇÃO (layout v9)
# ═══════════════════════════════════════════════════════════════════
cfg1,cfg2,cfg3,cfg4,cfg5=st.tabs([
    "⚙️ Parâmetros","📐 Fórmulas","🏦 Rating & PE","💰 Funding por faixas","🔆 Modo guiado"
])

# ── ABA 1: PARÂMETROS ─────────────────────────────────────────────
with cfg1:
    c1,c2,c3,c4=st.columns(4)
    with c1:
        st.markdown('<div class="sec-hdr">Operação base</div>',unsafe_allow_html=True)
        C=st.number_input("Capital (R$)",1e3,5e7,200_000.0,1e3,format="%.2f")
        T=st.number_input("Prazo (meses)",1,360,60,1)
        tc=st.number_input("Carência (meses)",0,T,12,1)
        if tc>0:
            pagar_car=st.radio("Na carência",["Capitaliza","Paga corrente"],horizontal=True,label_visibility="collapsed")=="Paga corrente"
        else: pagar_car=False
        exig_nome=st.selectbox("Exig. juros",list(EXIG_OPTS.keys()),label_visibility="collapsed")
        z_J=EXIG_OPTS[exig_nome]
        lamb=st.number_input("λ Margem alvo % a.m.",0.0,5.0,float(receita_base["params_margem"]["lamb_am"]),0.01,format="%.4f",
                              help="No Modo 1: λ define a taxa mínima. No Modo 2: verificação de viabilidade.")

    with c2:
        st.markdown('<div class="sec-hdr">Taxa ativa (Modo 2 / verificação)</div>',unsafe_allow_html=True)
        if st.session_state.modo_app=="modo2":
            st.caption("No Modo 2 a taxa ao cliente é calculada automaticamente pelo custo+spread de cada faixa.")
            i_am=float(op.get("i_am",1.43))
        else:
            usar_merc=st.checkbox("Comparar com taxa de mercado",help="Verifica se a taxa de mercado cobre a taxa mínima calculada")
            if usar_merc:
                ct1,ct2=st.columns(2)
                with ct1: i_am_fixo=st.number_input("Mercado %a.m.",0.01,20.0,1.20,0.01,format="%.4f",key="txf")
                with ct2: pct_cdi_v=st.number_input("% CDI",10.0,500.0,111.0,1.0,format="%.1f",key="txc")
                modo_tx=st.radio("Ref.",["Fixa","% CDI"],horizontal=True,label_visibility="collapsed",key="txr")
                i_am=i_am_fixo if modo_tx=="Fixa" else taxa_de_pct_cdi(pct_cdi_v,float(op.get("cdi_am",1.08)))
                st.session_state["taxa_merc_m1"]=i_am
            else:
                i_am=float(op.get("i_am",1.43))
                st.session_state["taxa_merc_m1"]=None
            cdi_am_input=st.number_input("CDI ref. % a.m.",0.01,10.0,float(op.get("cdi_am",1.08)),0.01,format="%.4f",
                                          help="Usado para exibir taxa mínima em % CDI")
            st.session_state["cdi_am_m1"]=cdi_am_input

        st.markdown('<div class="sec-hdr">Modo 2 — método de composição</div>',unsafe_allow_html=True)
        if st.session_state.modo_app=="modo2":
            metodo_m2=st.radio("Composição da taxa",["Produtória: (1+custo)×(1+spread)−1","Soma: custo+spread"],
                                label_visibility="collapsed",key="m2met")
            usar_prod=(metodo_m2.startswith("P"))
            st.session_state["usar_prod_m2"]=usar_prod
            cdi_am_input=float(op.get("cdi_am",1.08))
            st.session_state["cdi_am_m1"]=cdi_am_input

    with c3:
        st.markdown('<div class="sec-hdr">Custos e tributos</div>',unsafe_allow_html=True)
        custo_cont=st.number_input("Custo contrat. R$",0.0,1e5,float(cst.get("custo_cont",0)),100.0)
        custo_manut=st.number_input("Manut./parc R$",0.0,500.0,float(cst.get("custo_manut",0)),1.0)
        alfa_pc=st.number_input("PASEP+COF %",0.0,10.0,float(cst.get("alfa_pc",4.65)),0.01)
        alfa_iss=st.number_input("ISS %",0.0,10.0,float(cst.get("alfa_iss",5.0)),0.1)
        alfa_ir=st.number_input("IRPJ+CSLL %",0.0,60.0,float(cst.get("alfa_ir_cs",45.0)),1.0)
        st.markdown('<div class="sec-hdr">Basileia III</div>',unsafe_allow_html=True)
        FPR=st.slider("FPR %",0.0,150.0,float(rsc.get("FPR",75.0)),5.0)
        K=st.slider("K %",5.0,20.0,float(rsc.get("K",11.0)),0.25)
        Kp=st.slider("Kp %",5.0,20.0,float(rsc.get("Kp",9.75)),0.25)
        F_aval=st.number_input("F",1.0,30.0,float(rsc.get("F",10.95)),0.05)
        FCC=st.slider("FCC %",0.0,100.0,float(rsc.get("FCC",20.0)),5.0)

    with c4:
        st.markdown('<div class="sec-hdr">Provisão e PE</div>',unsafe_allow_html=True)
        fase1_am=st.number_input("IPP Fase1 %a.m.",0.0,1.0,float(rsc.get("fase1_am",0.175)),0.005,format="%.4f",
                                   help="Provisão regulatória de base. A PE do rating é usada na simulação por rating.")
        usa_f2=st.checkbox("Fase2 IPP",value=bool(rsc.get("usa_fase2",False)))
        fase2_pct=st.number_input("Fase2 %",0.0,10.0,float(rsc.get("fase2_pct",1.5)),0.1) if usa_f2 else 0.0
        tem_pe=st.checkbox("PE/Balloon")
        pe_pct=st.slider("PE %",0.0,30.0,0.0,0.5) if tem_pe else 0.0
        pe_parc=st.number_input("Parcela PE",1,T,T) if tem_pe else None
        with st.expander("Curvas DIFIN (visualização)"):
            arq=st.file_uploader("Upload .txt",type=["txt"],label_visibility="collapsed")
            if arq:
                try: st.session_state.curvas=parse_difin(arq.read().decode("utf-8","ignore")); st.success(f"✓ {len(st.session_state.curvas)} curvas")
                except Exception as e: st.error(str(e))
            if st.session_state.curvas:
                tv=st.selectbox("Ver",sorted(st.session_state.curvas.keys()),
                   format_func=lambda x:f"{x:02d}—{TIPOS_CURVA.get(x,'?')}",label_visibility="collapsed")
                if tv in st.session_state.curvas:
                    c_=st.session_state.curvas[tv]
                    st.line_chart(pd.DataFrame({"DU":[n.du for n in c_.nos],"Taxa":[n.taxa for n in c_.nos]}).set_index("DU"),color=["#1F3B6E"])

# ── ABA 2: FÓRMULAS ───────────────────────────────────────────────
with cfg2:
    for mod_key,mod in SLOTS.items():
        cor=mod["cor"]; bg=mod["bg"]
        st.markdown(f"<div class='mod-hdr' style='background:{cor}'>{mod['label']}</div><div class='mod-body'>",unsafe_allow_html=True)
        slots_items=list(mod["slots"].items()); ncols_mod=min(len(slots_items),4)
        cols_mod=st.columns(ncols_mod)
        for i_sl,(sn,slot) in enumerate(slots_items):
            with cols_mod[i_sl%ncols_mod]:
                render_slot(sn,slot,cor,bg)
        st.markdown("</div>",unsafe_allow_html=True)
    bdg="".join(f"<span class='badge'>{ch.split('.')[-1]}:{cod}</span>" for ch,cod in st.session_state.slots_sel.items())
    st.markdown(bdg,unsafe_allow_html=True)
    if st.button("↺ Restaurar defaults",use_container_width=True):
        st.session_state.slots_sel=dict(DEFAULTS); st.rerun()

# ── ABA 3: RATING & PE ────────────────────────────────────────────
with cfg3:
    cr1,cr2,cr3=st.columns([1.2,1.5,1.3])
    with cr1:
        st.markdown('<div class="sec-hdr">Arquivo de PE por rating</div>',unsafe_allow_html=True)
        cdl1,cdl2=st.columns(2)
        with cdl1: st.download_button("⬇ Excel",_gerar_template_excel(),"template_PE.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
        with cdl2: st.download_button("⬇ CSV",gerar_csv_template(),"template_PE.csv","text/csv",use_container_width=True)
        arq_pe=st.file_uploader("Upload PE",type=["xlsx","xls","csv"],label_visibility="collapsed")
        if arq_pe:
            try:
                pd_dict,avisos=carregar_pe_arquivo(arq_pe.read(),arq_pe.name)
                st.session_state.pd_upload=pd_dict; st.success(f"✓ {len(pd_dict)} ratings")
                for av in avisos: st.warning(av)
            except Exception as e: st.error(str(e))
        pd_ativo=st.session_state.pd_upload or PD_DEFAULTS_AA
        st.caption(f"PD ativo: {'arquivo' if st.session_state.pd_upload else 'defaults'}")
        lgd=st.number_input("LGD %",0.0,100.0,float(LGD_DEFAULT),1.0)
        rar_min=st.number_input("RAR_G mínimo %",0.0,30.0,8.0,0.5,format="%.1f")
        rsple_min=st.number_input("RSPLE mínimo %",0.0,30.0,6.0,0.5,format="%.1f")
        criterios=CriteriosViabilidade(rar_g_min=rar_min,rsple_min=rsple_min)
    with cr2:
        st.markdown('<div class="sec-hdr">Ratings a calcular</div>',unsafe_allow_html=True)
        ratings_sel=st.multiselect("Ratings",options=RATINGS_ORDEM,
            default=["A01","A03","A05","B01","B03","B05","C01","C03","D01","E01"],
            label_visibility="collapsed")
        tab_preview=build_tabela(pd_ativo)
        letras_u=list(dict.fromkeys(r[0] for r in RATINGS_ORDEM))
        st.dataframe(pd.DataFrame([{"Letra":l,
            "PD min-max":f"{min(tab_preview[r].pd_aa for r in RATINGS_ORDEM if r.startswith(l)):.2f}–{max(tab_preview[r].pd_aa for r in RATINGS_ORDEM if r.startswith(l)):.2f}%",
            "PE médio %am":f"{sum(tab_preview[r].pe_am for r in RATINGS_ORDEM if r.startswith(l))/sum(1 for r in RATINGS_ORDEM if r.startswith(l)):.4f}%"}
            for l in letras_u]),hide_index=True,use_container_width=True,height=270)
    with cr3:
        st.markdown('<div class="sec-hdr">Modo 1 — taxa de mercado</div>',unsafe_allow_html=True)
        st.caption("Informe na aba Parâmetros (coluna Taxa ativa). Aqui são só os botões de execução.")
        st.markdown("---")
        btn_calc=st.button("⚡ Calcular taxa mínima por rating",type="primary",use_container_width=True,
                           disabled=(len(ratings_sel)==0),key="btn_calc_m1")
        st.markdown("---")
        st.markdown('<div class="sec-hdr">Modo 2 — simular faixas</div>',unsafe_allow_html=True)
        btn_m2=st.button("🔄 Simular faixas × ratings",type="primary",use_container_width=True,
                          disabled=(len(ratings_sel)==0 or st.session_state.modo_app!="modo2"),key="btn_m2")

# ── ABA 4: FUNDING POR RATING ─────────────────────────────────────
with cfg4:
    st.markdown("**Custo de funding específico por rating**")
    st.caption("Cada rating tem seu próprio custo de captação. O cálculo usa exatamente esse custo para aquele rating — sem médias por letra.")

    sub1,sub2=st.tabs(["📥 Manual / Arquivo","📊 Resumo por letra"])

    with sub1:
        f4a,f4b=st.columns([1.5,1.2])
        with f4a:
            st.markdown('<div class="sec-hdr">Entrada de dados</div>',unsafe_allow_html=True)
            modo_fund=st.radio("Modo de entrada",["Manual por rating","Upload de arquivo"],
                               horizontal=True,label_visibility="collapsed",key="mfund")

            if modo_fund=="Upload de arquivo":
                cdl1,cdl2=st.columns(2)
                with cdl1: st.download_button("⬇ Template Excel",gerar_xlsx_funding(),"template_funding_rating.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
                with cdl2: st.download_button("⬇ Template CSV",gerar_csv_funding(),"template_funding_rating.csv",
                    "text/csv",use_container_width=True)
                arq_fund=st.file_uploader("Upload custo de funding (.xlsx/.csv)",
                                          type=["xlsx","xls","csv"],label_visibility="collapsed",key="arqfund")
                if arq_fund:
                    try:
                        fund_dict,avisos_f=carregar_funding_arquivo(arq_fund.read(),arq_fund.name)
                        st.session_state.funding_por_rating=fund_dict
                        st.success(f"✓ {len(fund_dict)} ratings carregados")
                        for av in avisos_f: st.warning(av)
                    except Exception as e: st.error(str(e))
            else:
                # Manual: editor por letra com expanders
                letras=list(dict.fromkeys(r[0] for r in RATINGS_ORDEM))
                for letra in letras:
                    cods=[r for r in RATINGS_ORDEM if r.startswith(letra)]
                    with st.expander(f"Letra {letra} — {cods[0]} a {cods[-1]}", expanded=(letra in ["A","B"])):
                        n=len(cods); ncols=min(n,5); cols_l=st.columns(ncols)
                        for i_c,cod in enumerate(cods):
                            with cols_l[i_c%ncols]:
                                val_atual=st.session_state.funding_por_rating.get(cod,FUNDING_DEFAULTS_AA.get(cod,8.0))
                                novo=st.number_input(cod,0.0,40.0,float(val_atual),0.25,
                                                     format="%.2f",key=f"fr_{cod}",
                                                     label_visibility="visible")
                                st.session_state.funding_por_rating[cod]=novo

            col_r1,col_r2=st.columns(2)
            with col_r1:
                if st.button("⟳ Restaurar defaults",use_container_width=True,key="freset"):
                    st.session_state.funding_por_rating=dict(FUNDING_DEFAULTS_AA); st.rerun()
            with col_r2:
                if st.button("= Mesmo valor para todos",use_container_width=True,key="fmesmo"):
                    st.session_state["fund_uniforme"]=True

            if st.session_state.get("fund_uniforme"):
                val_uni=st.number_input("Custo único % a.a. para todos os ratings",0.0,40.0,8.0,0.25,format="%.2f",key="funi")
                if st.button("Aplicar",use_container_width=True,key="fapl"):
                    for cod in RATINGS_ORDEM:
                        st.session_state.funding_por_rating[cod]=val_uni
                    st.session_state["fund_uniforme"]=False; st.rerun()

        with f4b:
            st.markdown('<div class="sec-hdr">Custo atual por rating</div>',unsafe_allow_html=True)
            fund_atual=st.session_state.funding_por_rating
            rows_fa=[]
            for cod in RATINGS_ORDEM:
                v=fund_atual.get(cod,FUNDING_DEFAULTS_AA.get(cod,8.0))
                rows_fa.append({"Rating":cod,"Custo % a.a.":f"{v:.2f}%",
                                 "Custo % a.m.":f"{taxa_am_de_aa(v):.4f}%"})
            st.dataframe(pd.DataFrame(rows_fa),hide_index=True,use_container_width=True,height=520)

    with sub2:
        st.markdown("**Resumo por letra — visão consolidada**")
        fund_atual=st.session_state.funding_por_rating
        rows_letra=resumo_por_letra(fund_atual)
        df_letra=pd.DataFrame(rows_letra)
        CORES_LETRA={"A":"#E8F5ED","B":"#EEF2FA","C":"#F3EFFE",
                     "D":"#FFF3E0","E":"#FCEBEB","F":"#FFD9D9","G":"#FFCCCC"}
        def hl_letra(row):
            bg=CORES_LETRA.get(row.get("Letra",""),"")
            return [f"background:{bg}"]*len(row) if bg else [""]*len(row)
        st.dataframe(df_letra.style.apply(hl_letra,axis=1),hide_index=True,use_container_width=True)
        st.caption("O cálculo por rating usa o custo individual de cada rating — esta tabela é apenas para visualização.")

    # WACC para Modo 2 (mantém compatibilidade)
    st.markdown("---")
    st.markdown('<div class="sec-hdr">Funding por faixas — Modo 2 (spread travado)</div>',unsafe_allow_html=True)
    cf1m2,cf2m2=st.columns([2.5,1.2])
    with cf1m2:
        if st.button("➕ Adicionar faixa",key="f4add"):
            st.session_state.faixas_funding.append({"nome":f"Faixa {len(st.session_state.faixas_funding)+1}","taxa_if_aa":4.5,"taxa_fs_aa":2.0,"participacao":0.0}); st.rerun()
        rm_list=[]
        for i_f,faixa in enumerate(st.session_state.faixas_funding):
            fa,fb,fc,fd,frm=st.columns([2.2,0.8,0.8,0.8,0.3])
            with fa: faixa["nome"]=st.text_input("",faixa["nome"],key=f"fn_{i_f}",label_visibility="collapsed")
            with fb: faixa["taxa_fs_aa"]=st.number_input("Custo%aa",0.0,30.0,float(faixa.get("taxa_fs_aa",2.0)),0.1,format="%.2f",key=f"fsc_{i_f}")
            with fc: faixa["taxa_if_aa"]=st.number_input("Spread%aa",0.0,20.0,float(faixa.get("taxa_if_aa",4.5)),0.1,format="%.2f",key=f"ftif_{i_f}")
            with fd: faixa["participacao"]=st.number_input("Part%",0.0,100.0,float(faixa.get("participacao",25.0)),5.0,format="%.1f",key=f"fpart_{i_f}")
            with frm:
                if st.button("✕",key=f"frm_{i_f}"): rm_list.append(i_f)
        for i in sorted(rm_list,reverse=True): st.session_state.faixas_funding.pop(i); st.rerun()
        if st.button("⟳ Preset padrão",key="fpre"):
            st.session_state.faixas_funding=copy.deepcopy(FAIXAS_PADRAO); st.rerun()
    with cf2m2:
        w=wacc_faixas(st.session_state.faixas_funding)
        ok_w=abs(w["soma_part"]-100)<0.5
        if not ok_w: st.warning(f"Soma = {w['soma_part']:.1f}%")
        else: st.success(f"✓ Soma = {w['soma_part']:.1f}%")
        st.metric("WACC % a.a.",f"{w['wacc_aa']:.4f}%")
        st.metric("WACC % a.m.",f"{w['wacc_am']:.6f}%")
        usar_wacc=st.checkbox("Usar WACC como custo base",value=True,key="fwacc")
        cdi_ef=w["wacc_am"] if (usar_wacc and ok_w) else float(op.get("cdi_am",1.08))
        if not ok_w: cdi_ef=float(op.get("cdi_am",1.08))

# ── ABA 5: MODO GUIADO ────────────────────────────────────────────
with cfg5:
    st.markdown("**Checklist de preenchimento**")
    STEPS_INFO=[
        ("Produto","Produto selecionado no cabeçalho","⚙️ Parâmetros"),
        ("Operação","Capital, prazo, carência, exig. juros","⚙️ Parâmetros"),
        ("λ Margem","Margem alvo definida","⚙️ Parâmetros"),
        ("Funding","Faixas e WACC configurados","💰 Funding"),
        ("Custos","PASEP/COF, ISS, IR/CS, custos admin.","⚙️ Parâmetros"),
        ("Basileia","FPR, K, Kp, F, FCC","⚙️ Parâmetros"),
        ("PE / IPP","Fase1, Fase2, LGD","⚙️ Parâmetros / 🏦 Rating"),
        ("Fórmulas","Verificar slots Z/F de cada módulo","📐 Fórmulas"),
        ("Ratings","Selecionar ratings para calcular","🏦 Rating & PE"),
    ]
    cols_steps=st.columns(3)
    for i,(nome,desc,aba) in enumerate(STEPS_INFO):
        with cols_steps[i%3]:
            st.markdown(
                f"<div style='border:1px solid var(--color-border-tertiary);border-radius:7px;padding:8px;margin-bottom:6px'>"
                f"<div style='font-size:11px;font-weight:500;color:var(--color-text-primary)'>{i+1}. {nome}</div>"
                f"<div style='font-size:10px;color:var(--color-text-secondary);margin-top:2px'>{desc}</div>"
                f"<div style='font-size:9px;color:#1F3B6E;margin-top:3px'>→ aba {aba}</div>"
                f"</div>",unsafe_allow_html=True)

st.markdown("---")

# ── Monta receita base ────────────────────────────────────────────
cdi_am_ref=st.session_state.get("cdi_am_m1",float(op.get("cdi_am",1.08)))
rc_base=_build_rc(produto_sel,C,T,tc,i_am,cdi_ef,lamb,pagar_car,z_J,
                  custo_cont,custo_manut,alfa_pc,alfa_iss,alfa_ir,
                  FPR,K,Kp,F_aval,FCC,fase1_am,usa_f2,fase2_pct,tem_pe,pe_pct,pe_parc)

# ── Execuções ─────────────────────────────────────────────────────
if btn_calc:
    with st.spinner(f"Calculando taxa mínima para {len(ratings_sel)} ratings..."):
        tab_rat=build_tabela(pd_ativo,lgd); resultados=[]; prog=st.progress(0)
        fund_am_por_rat=funding_am_por_rating(st.session_state.funding_por_rating)
        for i_r,cod in enumerate(ratings_sel):
            if cod in tab_rat:
                fund_am_rat = fund_am_por_rat.get(cod)
                r=calcular_taxa_min(rc_base,tab_rat[cod],rar_min,rsple_min,cdi_am_ref,fund_am_rat)
                if r:
                    r["custo_funding_aa"] = st.session_state.funding_por_rating.get(cod, FUNDING_DEFAULTS_AA.get(cod,8.0))
                    r["custo_funding_am"] = fund_am_rat or cdi_am_ref
                    resultados.append(r)
            prog.progress((i_r+1)/len(ratings_sel))
        prog.empty()
        st.session_state.resultado_m1={"resultados":resultados,"modo":"modo1",
            "taxa_merc":st.session_state.get("taxa_merc_m1"),"cdi_am":cdi_am_ref}

if btn_m2:
    with st.spinner("Simulando faixas × ratings..."):
        tab_rat2=build_tabela(pd_ativo,lgd); todas=[]; usar_prod=st.session_state.get("usar_prod_m2",True)
        for faixa in st.session_state.faixas_funding:
            custo=faixa["taxa_fs_aa"]; spread=faixa["taxa_if_aa"]
            if custo is None: custo=2.0
            if spread is None: spread=4.5
            tx_am=(taxa_produtoria_am if usar_prod else taxa_soma_am)(custo,spread)
            custo_am=taxa_am_de_aa(custo)
            rc_f=_build_rc(produto_sel,C,T,tc,tx_am,custo_am,lamb,pagar_car,z_J,
                           custo_cont,custo_manut,alfa_pc,alfa_iss,alfa_ir,
                           FPR,K,Kp,F_aval,FCC,fase1_am,usa_f2,fase2_pct,tem_pe,pe_pct,pe_parc)
            txs_f={cod:tx_am for cod in ratings_sel}
            try:
                sim_f=simular_portfolio_v2(rc_f,tab_rat2,criterios,ratings_sel,txs_f)
                tx_aa=((1+custo/100)*(1+spread/100)-1)*100 if usar_prod else (custo+spread)
                for l in sim_f["linhas"]:
                    l.update({"Faixa":faixa["nome"][:25],"Custo %aa":round(custo,2),
                              "Spread %aa":round(spread,2),"Taxa cliente %aa":round(tx_aa,4),
                              "Método":"Prod." if usar_prod else "Soma"})
                    todas.append(l)
            except Exception as e: st.error(str(e))
        st.session_state.resultado_m1={"modo2_linhas":todas,"modo":"modo2"}

# ═══════════════════════════════════════════════════════════════════
# PARTE INFERIOR — ABAS DE RESULTADO (layout v9)
# ═══════════════════════════════════════════════════════════════════
if not st.session_state.resultado_m1:
    st.info("Configure os parâmetros nas abas acima e clique em **⚡ Calcular taxa mínima** ou **🔄 Simular faixas**")
    st.stop()

rm=st.session_state.resultado_m1

r1,r2,r3,r4,r5,r6=st.tabs([
    "📊 Taxa mínima","🧩 Decomposição","🎯 vs Mercado",
    "📋 Planilha","🏦 Faixas (Modo 2)","⚖️ Histórico"
])

# ── Tab 1: Taxa mínima ────────────────────────────────────────────
with r1:
    if rm.get("modo")!="modo1":
        st.info("Execute o Modo 1 para ver esta aba.")
    else:
        resultados=rm["resultados"]; cdi_ref=rm["cdi_am"]; taxa_merc=rm.get("taxa_merc")
        fronteira=None
        for r in resultados:
            if r.get("taxa_min_am"): fronteira=r["rating"]

        if fronteira:
            r_f=next(r for r in resultados if r["rating"]==fronteira)
            cor_f=CORES_RATING.get(fronteira[0],"#333"); bg_f=BACKCORES_RATING.get(fronteira[0],"#F5F5F5")
            st.markdown(
                f"<div class='front-banner' style='background:{bg_f};border:2px solid {cor_f}'>"
                f"<div style='font-size:30px;font-weight:500;color:{cor_f};min-width:55px'>{fronteira}</div>"
                f"<div style='flex:1'>"
                f"<div style='font-size:10px;font-weight:700;color:{cor_f};text-transform:uppercase;letter-spacing:.06em'>Fronteira de viabilidade</div>"
                f"<div style='font-size:12px;color:{cor_f};margin-top:3px'>"
                f"{sum(1 for r in resultados if r.get('taxa_min_am'))} de {len(resultados)} ratings viáveis | "
                f"RAR_G≥{rar_min:.1f}% | RSPLE≥{rsple_min:.1f}% | λ={lamb:.4f}%a.m.</div></div>"
                f"<div style='background:#fff;border:0.5px solid {cor_f};border-radius:7px;padding:6px 10px;text-align:center'>"
                f"<div style='font-size:16px;font-weight:500;color:{cor_f}'>{r_f['taxa_min_am']:.4f}%</div>"
                f"<div style='font-size:9px;color:{cor_f}'>Taxa mín. {fronteira}</div></div>"
                f"<div style='background:#fff;border:0.5px solid {cor_f};border-radius:7px;padding:6px 10px;text-align:center'>"
                f"<div style='font-size:16px;font-weight:500;color:{cor_f}'>{r_f['pct_cdi']:.1f}%</div>"
                f"<div style='font-size:9px;color:{cor_f}'>% CDI</div></div>"
                f"<div style='background:#fff;border:0.5px solid {cor_f};border-radius:7px;padding:6px 10px;text-align:center'>"
                f"<div style='font-size:16px;font-weight:500;color:{cor_f}'>{r_f['RAR_G']:.2f}%</div>"
                f"<div style='font-size:9px;color:{cor_f}'>RAR_G</div></div>"
                f"<div style='background:#fff;border:0.5px solid {cor_f};border-radius:7px;padding:6px 10px;text-align:center'>"
                f"<div style='font-size:16px;font-weight:500;color:{cor_f}'>{r_f['RSPLE']:.2f}%</div>"
                f"<div style='font-size:9px;color:{cor_f}'>RSPLE</div></div>"
                f"</div>",unsafe_allow_html=True)
        else:
            st.error("Nenhum rating viável com os parâmetros configurados.")

        rows_min=[]
        for r in resultados:
            tx=r.get("taxa_min_am")
            merc_s=""
            if taxa_merc and tx:
                d=taxa_merc-tx
                merc_s=f"+{d:.4f}%" if d>=0 else f"{d:.4f}%"
            rows_min.append({
                "Rating":r["rating"],"PD % a.a.":f"{r['pd_aa']:.2f}%","PE % a.m.":f"{r['pe_am']:.4f}%",
                "Fund. % a.a.":f"{r.get('custo_funding_aa',0):.2f}%" if tx else "—",
                "Fund. % a.m.":f"{r.get('custo_funding_am',0):.4f}%" if tx else "—",
                "Taxa mín % a.m.":f"{tx:.4f}%" if tx else "Inviável",
                "Taxa mín % a.a.":f"{r['taxa_min_aa']:.4f}%" if tx else "—",
                "% CDI ref":f"{r['pct_cdi']:.1f}%" if tx else "—",
                "RSPLE %":f"{r['RSPLE']:.4f}%","RAR_G %":f"{r['RAR_G']:.4f}%",
                "Duration":f"{r.get('Duration',0):.1f}m" if tx else "—",
                "TIR % a.a.":f"{r['TIR']:.4f}%" if r.get("TIR") else "—",
                **({"vs Mercado":merc_s} if taxa_merc else {}),
                "Status":"✅ Viável" if tx else "❌ Inviável",
            })
        df_min=pd.DataFrame(rows_min)
        def hl_min(row):
            s=str(row.get("Status","")); r_=row.get("Rating","")
            if fronteira and r_==fronteira: return ["background:#D6E4F5;font-weight:bold"]*len(row)
            if "✅" in s: return ["background:#E8F5ED"]*len(row)
            return ["background:#FCEBEB"]*len(row)
        st.dataframe(df_min.style.apply(hl_min,axis=1),hide_index=True,use_container_width=True)
        if taxa_merc:
            st.caption(f"Taxa de mercado: **{taxa_merc:.4f}% a.m.** ({pct_cdi_de_taxa(taxa_merc,cdi_ref):.1f}% CDI)")

# ── Tab 2: Decomposição ───────────────────────────────────────────
with r2:
    if rm.get("modo")!="modo1":
        st.info("Execute o Modo 1.")
    else:
        resultados=rm["resultados"]; fronteira=None
        for r in resultados:
            if r.get("taxa_min_am"): fronteira=r["rating"]
        viav=[r for r in resultados if r.get("taxa_min_am")]
        st.markdown("**Decomposição da taxa mínima — quanto de cada componente na receita de JA**")
        if viav:
            for r in viav:
                d=r["decomp"]; tot=d['f']+d['p']+d['t']+d['m'] or 1
                pf,pp,pt,pm=d['f']/tot*100,d['p']/tot*100,d['t']/tot*100,d['m']/tot*100
                front_lbl=" ← fronteira" if r["rating"]==fronteira else ""
                st.markdown(
                    f"<div style='display:flex;align-items:center;gap:10px;margin-bottom:7px'>"
                    f"<span style='font-size:11px;font-weight:500;width:36px;color:var(--color-text-primary)'>{r['rating']}</span>"
                    f"<span style='font-size:10px;color:var(--color-text-secondary);width:65px'>{r['taxa_min_am']:.4f}%am</span>"
                    f"<div style='flex:1'><div class='decomp-bar'>"
                    f"<div class='seg-f' style='width:{pf:.1f}%'></div>"
                    f"<div class='seg-p' style='width:{pp:.1f}%'></div>"
                    f"<div class='seg-t' style='width:{pt:.1f}%'></div>"
                    f"<div class='seg-m' style='width:{pm:.1f}%'></div>"
                    f"</div></div>"
                    f"<span style='font-size:10px;color:var(--color-text-secondary);width:200px'>"
                    f"F:{d['f']}% P:{d['p']}% T:{d['t']}% M:{d['m']}%{front_lbl}</span>"
                    f"</div>",unsafe_allow_html=True)
            st.markdown(
                "<div style='display:flex;gap:14px;margin-top:8px'>"
                "<div style='display:flex;align-items:center;gap:4px;font-size:11px'><div style='width:10px;height:10px;border-radius:2px;background:#378ADD'></div>Funding</div>"
                "<div style='display:flex;align-items:center;gap:4px;font-size:11px'><div style='width:10px;height:10px;border-radius:2px;background:#E24B4A'></div>PE / Risco</div>"
                "<div style='display:flex;align-items:center;gap:4px;font-size:11px'><div style='width:10px;height:10px;border-radius:2px;background:#EF9F27'></div>Tributos</div>"
                "<div style='display:flex;align-items:center;gap:4px;font-size:11px'><div style='width:10px;height:10px;border-radius:2px;background:#3B6D11'></div>Margem</div>"
                "</div>",unsafe_allow_html=True)
        else: st.info("Nenhum rating viável.")

# ── Tab 3: vs Mercado ─────────────────────────────────────────────
with r3:
    if rm.get("modo")!="modo1":
        st.info("Execute o Modo 1.")
    else:
        resultados=rm["resultados"]; cdi_ref=rm["cdi_am"]; taxa_merc=rm.get("taxa_merc")
        if not taxa_merc:
            st.info("Marque 'Comparar com taxa de mercado' na aba Parâmetros e recalcule.")
        else:
            st.markdown(f"**Taxa de mercado:** {taxa_merc:.4f}% a.m. = {pct_cdi_de_taxa(taxa_merc,cdi_ref):.1f}% CDI")
            rows_vs=[]
            for r in resultados:
                tx=r.get("taxa_min_am")
                if not tx: continue
                delta=taxa_merc-tx; cobre=(delta>=0)
                rows_vs.append({"Rating":r["rating"],"Taxa mín %am":f"{tx:.4f}%",
                    "Mercado %am":f"{taxa_merc:.4f}%","Folga %am":f"{delta:+.4f}%",
                    "Folga pp CDI":f"{pct_cdi_de_taxa(taxa_merc,cdi_ref)-r['pct_cdi']:+.1f}pp",
                    "Status":"✅ Cobre" if cobre else "❌ Não cobre"})
            df_vs=pd.DataFrame(rows_vs)
            def hl_vs(row):
                if "✅" in str(row.get("Status","")): return ["background:#E8F5ED"]*len(row)
                return ["background:#FCEBEB"]*len(row)
            st.dataframe(df_vs.style.apply(hl_vs,axis=1),hide_index=True,use_container_width=True)
            vals=[r for r in resultados if r.get("taxa_min_am")]
            if vals:
                df_g=pd.DataFrame({"Rating":[r["rating"] for r in vals],
                                    "Taxa mínima":[r["taxa_min_am"] for r in vals],
                                    "Mercado":[taxa_merc]*len(vals)}).set_index("Rating")
                st.line_chart(df_g,color=["#C75B00","#1F3B6E"])
                st.caption("🟠 Taxa mínima por rating — 🔵 Taxa de mercado | abaixo da linha laranja = não cobre")

# ── Tab 4: Planilha (Excel) ───────────────────────────────────────
with r4:
    if rm.get("modo")=="modo1":
        resultados=rm.get("resultados",[])
        viav=[r for r in resultados if r.get("taxa_min_am") and "_df" in r]
        fronteira_r=None
        for r in resultados:
            if r.get("taxa_min_am"): fronteira_r=r["rating"]

        if not viav:
            st.info("Nenhum rating viável calculado ainda.")
        else:
            st.markdown("**Fluxo completo por período — nomenclatura do manual DIFIN/GEAFI**")
            opcoes=[r["rating"] for r in viav]
            rat_sel=st.selectbox("Rating para exibir o fluxo:",opcoes,
                                  index=opcoes.index(fronteira_r) if fronteira_r in opcoes else 0)
            r_sel=next(r for r in viav if r["rating"]==rat_sel)
            st.caption(
                f"Taxa mínima: **{r_sel['taxa_min_am']:.4f}% a.m.** = **{r_sel['pct_cdi']:.1f}% CDI** | "
                f"RAR_G: **{r_sel['RAR_G']:.4f}%** | RSPLE: **{r_sel['RSPLE']:.4f}%** | "
                f"FLA VP: **{r_sel['FLA_VP']:,.2f}**")
            render_planilha(r_sel.get("_df"), tc, rat_sel)
            # Download Excel
            caminho_tmp=os.path.join(os.path.dirname(__file__),"_exp.xlsx")
            try:
                para_excel({"fluxo":r_sel["_df"],"resumo":r_sel["_rs"]},caminho_tmp)
                with open(caminho_tmp,"rb") as f_:
                    st.download_button(
                        f"⬇ Baixar Excel — {rat_sel} ({r_sel['taxa_min_am']:.4f}%am)",
                        f_.read(),f"Fluxo_{rat_sel}_{r_sel['taxa_min_am']:.4f}pct.xlsx",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True)
            except: pass
    elif rm.get("modo")=="modo2":
        st.info("Planilha disponível apenas no Modo 1. No Modo 2 os fluxos individuais não são armazenados.")
    else:
        st.info("Calcule primeiro.")

# ── Tab 5: Faixas (Modo 2) ────────────────────────────────────────
with r5:
    if rm.get("modo")!="modo2":
        st.info("Execute o Modo 2 para ver esta aba.")
    else:
        todas=rm["modo2_linhas"]
        faixas_u=list(dict.fromkeys(l.get("Faixa","") for l in todas))
        for fname in faixas_u:
            linhas_f=[l for l in todas if l.get("Faixa","")==fname]
            front_f=None
            for l in linhas_f:
                if l.get("Viável"): front_f=l["Rating"]
            tx_aa=linhas_f[0].get("Taxa cliente %aa",0) if linhas_f else 0
            custo=linhas_f[0].get("Custo %aa",0); spread=linhas_f[0].get("Spread %aa",0)
            met=linhas_f[0].get("Método",""); viav=sum(1 for l in linhas_f if l.get("Viável"))
            cor_f=CORES_RATING.get(front_f[0],"#888") if front_f else "#888"
            bg_f=BACKCORES_RATING.get(front_f[0],"#F5F5F5") if front_f else "#F5F5F5"
            st.markdown(
                f"<div style='background:{bg_f};border:1.5px solid {cor_f};border-radius:8px;padding:8px 14px;margin-bottom:8px'>"
                f"<b style='color:{cor_f}'>{fname}</b> &nbsp;|&nbsp; "
                f"Custo: {custo:.2f}%aa &nbsp;{met}&nbsp; Spread: {spread:.2f}%aa → "
                f"Taxa cliente: <b>{tx_aa:.4f}%aa</b> &nbsp;|&nbsp; "
                f"Fronteira: <b style='color:{cor_f}'>{front_f or 'Nenhum'}</b> &nbsp;|&nbsp; "
                f"{viav}/{len(linhas_f)} ratings viáveis"
                f"</div>",unsafe_allow_html=True)
            cols_show=["Rating","PD % a.a.","PE % a.m.","RSPLE %","RAR_G %","Total SP","Total MG","FLA VP","Viável"]
            rows_f=[{c:l.get(c,"—") for c in cols_show} for l in linhas_f]
            df_f=pd.DataFrame(rows_f)
            for col in ["RSPLE %","RAR_G %","PD % a.a.","PE % a.m."]:
                if col in df_f: df_f[col]=df_f[col].apply(lambda v:f"{v:.4f}" if isinstance(v,float) else str(v))
            for col in ["Total SP","Total MG"]:
                if col in df_f: df_f[col]=df_f[col].apply(lambda v:f"R$ {v:,.2f}" if isinstance(v,float) else str(v))
            df_f["FLA VP"]=df_f["FLA VP"].apply(lambda v:f"{v:,.2f}" if isinstance(v,float) else str(v))
            df_f["Viável"]=df_f["Viável"].apply(lambda v:"✅" if v is True else "❌" if v is False else str(v))
            def hlf(row):
                if "✅" in str(row.get("Viável","")): return ["background:#E8F5ED"]*len(row)
                if "❌" in str(row.get("Viável","")): return ["background:#FCEBEB"]*len(row)
                return [""]*len(row)
            st.dataframe(df_f.style.apply(hlf,axis=1),hide_index=True,use_container_width=True,height=260)
        st.download_button("⬇ CSV completo",
            pd.DataFrame([{c:l.get(c,"") for c in ["Faixa","Rating","Taxa cliente %aa","Custo %aa","Spread %aa","PD % a.a.","RSPLE %","RAR_G %","FLA VP","Viável"]} for l in todas]).to_csv(index=False,sep=";",decimal=",").encode(),
            "Portfolio_Modo2.csv","text/csv",use_container_width=True)

# ── Tab 6: Histórico ──────────────────────────────────────────────
with r6:
    if rm.get("modo")=="modo1":
        resultados=rm.get("resultados",[])
        fronteira_h=None
        for r in resultados:
            if r.get("taxa_min_am"): fronteira_h=r["rating"]
        nome_c=st.text_input("Nome deste cálculo",
            value=f"{produto_sel} — {fronteira_h or '—'} — λ={lamb:.4f}%am — WACC={cdi_ef:.4f}%am")
        if st.button("💾 Salvar no histórico",use_container_width=True):
            st.session_state.cenarios[nome_c]={
                "fronteira":fronteira_h,"modo":"Modo 1",
                "n_viaveis":sum(1 for r in resultados if r.get("taxa_min_am")),
                "rar_min":rar_min,"rsple_min":rsple_min,"lamb":lamb,"cdi":cdi_ef,
                "ratings":{r["rating"]:r.get("taxa_min_am") for r in resultados if r.get("taxa_min_am")},
            }; st.success("Salvo!")
    if st.session_state.cenarios:
        st.markdown("**Histórico:**")
        for nome,dados in st.session_state.cenarios.items():
            cor_c=CORES_RATING.get((dados.get("fronteira") or "A")[0],"#333")
            bg_c=BACKCORES_RATING.get((dados.get("fronteira") or "A")[0],"#F5F5F5")
            st.markdown(
                f"<div style='background:{bg_c};border:1px solid {cor_c};border-radius:7px;padding:8px 12px;margin-bottom:6px'>"
                f"<b style='color:{cor_c};font-size:12px'>{nome}</b><br>"
                f"<span style='font-size:11px;color:var(--color-text-secondary)'>"
                f"Fronteira: {dados.get('fronteira','—')} | {dados.get('n_viaveis',0)} ratings viáveis | "
                f"λ={dados.get('lamb',0):.4f}%am | WACC={dados.get('cdi',0):.4f}%am | {dados.get('modo','')}</span></div>",
                unsafe_allow_html=True)
        if st.button("🗑 Limpar histórico"): st.session_state.cenarios={}; st.rerun()
