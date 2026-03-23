"""
exporters.py
════════════
Exporta o resultado do engine_core para Excel (e CSV opcionalmente).
Abas: Parâmetros | Fluxo Financeiro | Indicadores
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Paleta de cores ───────────────────────────────────────────────────────────
AZ  = "1F3B6E";  AZ2 = "2E5FAB"; AZ_CL = "D6E4F5"
VD  = "1A5C3A";  VD_CL = "E8F5ED"
AM  = "C75B00";  AM_CL = "FFF3E0"
RX  = "4A1080"
CZ  = "F5F5F5";  BR  = "FFFFFF"

def _b():
    s = Side(border_style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _h(cel, txt, bg=AZ, fg=BR, sz=9, bold=True, wrap=True, ha="center"):
    cel.value = txt
    cel.font = Font(name="Calibri", bold=bold, size=sz, color=fg)
    cel.fill = PatternFill("solid", fgColor=bg)
    cel.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=wrap)
    cel.border = _b()

def _d(cel, val, fmt=None, bg=BR, fg="000000", bold=False, ha="right"):
    cel.value = val
    cel.font = Font(name="Calibri", size=9, color=fg, bold=bold)
    cel.fill = PatternFill("solid", fgColor=bg)
    cel.alignment = Alignment(horizontal=ha, vertical="center")
    cel.border = _b()
    if fmt: cel.number_format = fmt


def para_excel(resultado: dict, caminho: str) -> None:
    """Exporta resultado completo do engine para .xlsx."""
    wb = Workbook()
    wb.remove(wb.active)

    _aba_parametros(wb, resultado)
    _aba_fluxo(wb, resultado)
    _aba_indicadores(wb, resultado)

    wb.save(caminho)
    print(f"✓ Salvo: {caminho}")


def para_csv(resultado: dict, caminho: str) -> None:
    """Exporta o fluxo financeiro como CSV."""
    resultado["fluxo"].round(6).to_csv(caminho, index=False, sep=";", decimal=",")
    print(f"✓ CSV salvo: {caminho}")


# ══════════════════════════════════════════════════════════════════════════════
def _aba_parametros(wb, resultado):
    ws = wb.create_sheet("Parâmetros")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 18

    r  = resultado["resumo"]
    rc = resultado["receita"]

    # Título
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Crédito Direto ao Consumidor — {rc['nome']}"
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=BR)
    ws["A1"].fill = PatternFill("solid", fgColor=AZ)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24

    # Bloco: Dados da operação / Provisão de risco
    for col, lbl, cor in [("A","DADOS DA OPERAÇÃO", AZ2), ("D","PROVISÃO DE RISCO","7B3F00")]:
        c2 = "C" if col == "A" else "F"
        ws.merge_cells(f"{col}2:{c2}2")
        c = ws[col+"2"]
        c.value = lbl
        c.font = Font(name="Calibri", bold=True, size=10, color=BR)
        c.fill = PatternFill("solid", fgColor=cor)
        c.alignment = Alignment(horizontal="center")
        c.border = _b()
        for cc in (["B","C"] if col=="A" else ["E","F"]):
            ws[cc+"2"].fill = PatternFill("solid", fgColor=cor)
            ws[cc+"2"].border = _b()
    ws.row_dimensions[2].height = 16

    op_rows = [
        ("Valor Financiado", r["Capital (C)"], "R$"),
        ("Prazo (meses)", r["Prazo (T)"], None),
        ("Taxa ativa % a.m.", r["Taxa ativa % a.m."], "pct"),
        ("CDI % a.m.", r["CDI % a.m."], "pct"),
        ("λ Margem % a.m.", r["λ Margem % a.m."], "pct"),
        ("λ Margem % a.a.", r["λ Margem % a.a."], "pct"),
        ("PMTA Price", r.get("PMTA Price") or "—", "R$"),
        ("Fórmulas: JA/EBP/MG", f"{r.get('fórmula JA','—')} / {r.get('fórmula EBP','—')} / {r.get('fórmula MG','—')}", None),
    ]
    pr_rows = [
        ("FPR", r["FPR"], None),
        ("K regulatório", r["K regulatório"], None),
        ("K prudencial (Kp)", r["K prudencial (Kp)"], None),
        ("F / FCC", r["F / FCC"], None),
        ("CE regulatório (R$)", r["CE regulatório"], "R$"),
        ("CEIRB / CE gestão (R$)", r["CEIRB (gestão)"], "R$"),
        ("IPP Fase1 % a.m.", rc["params_risco"].get("fase1_am", "—"), "pct"),
        ("Atualiz. IPP", rc["params_risco"].get("atu_ipp", "—"), None),
    ]
    for i, ((al, av, af), (dl, dv, df_)) in enumerate(zip(op_rows, pr_rows), 3):
        bg = CZ if i%2==0 else BR
        _h(ws.cell(row=i, column=1), al, bg=bg, fg=AZ, bold=False, ha="left")
        fmta = "R$ #,##0.00" if af=="R$" else ("0.0000%" if af=="pct" else None)
        _d(ws.cell(row=i, column=2), av/100 if af=="pct" and isinstance(av,float) else av, fmta, bg=bg)
        ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor=bg)
        ws.cell(row=i, column=3).border = _b()
        _h(ws.cell(row=i, column=4), dl, bg=bg, fg=AZ, bold=False, ha="left")
        fmtd = "R$ #,##0.00" if df_=="R$" else ("0.0000%" if df_=="pct" else None)
        _d(ws.cell(row=i, column=5), dv/100 if df_=="pct" and isinstance(dv,float) else dv, fmtd, bg=bg)

    # Bloco: CUSTO DA OPERAÇÃO / PERFORMANCE
    ws.merge_cells("A11:C11")
    c = ws["A11"]; c.value="CUSTO DA OPERAÇÃO"
    c.font=Font(name="Calibri",bold=True,size=10,color=BR)
    c.fill=PatternFill("solid",fgColor=AM); c.alignment=Alignment(horizontal="center"); c.border=_b()
    for cc in ["B","C"]: ws[cc+"11"].fill=PatternFill("solid",fgColor=AM); ws[cc+"11"].border=_b()

    ws.merge_cells("D11:F11")
    c = ws["D11"]; c.value="PERFORMANCE"
    c.font=Font(name="Calibri",bold=True,size=10,color=BR)
    c.fill=PatternFill("solid",fgColor=RX); c.alignment=Alignment(horizontal="center"); c.border=_b()
    for cc in ["E","F"]: ws[cc+"11"].fill=PatternFill("solid",fgColor=RX); ws[cc+"11"].border=_b()

    cst = rc["params_custo"]
    custo_rows = [
        ("Custo contratação (R$)", cst.get("custo_cont",0), "R$"),
        ("Custo manutenção/período", cst.get("custo_manut",0), "R$"),
        ("PASEP+COFINS %", cst.get("alfa_pc",0), "pct"),
        ("ISS %", cst.get("alfa_iss",0), "pct"),
        ("IRPJ+CSLL %", cst.get("alfa_ir_cs",0), "pct"),
    ]
    perf_rows = [
        ("MG % a.a.", r["MG % a.a."], "pct"),
        ("RSPLE %", r["RSPLE %"], "pct"),
        ("RAR_G % (gestão, F1183)", r["RAR_G % (gestão, F1183)"], "pct"),
        ("Resultado PPS % a.m.", r.get("Resultado PPS % a.m.") or 0, "pct"),
        ("TIR % a.a.", r.get("TIR % a.a.") or 0, "pct"),
    ]
    for i, ((cl, cv, cf), (pl, pv, pf)) in enumerate(zip(custo_rows, perf_rows), 12):
        bg = CZ if i%2==0 else BR
        _h(ws.cell(row=i,column=1),cl,bg=bg,fg=AM,bold=False,ha="left")
        fmtc = "R$ #,##0.00" if cf=="R$" else "0.00%"
        _d(ws.cell(row=i,column=2), cv/100 if cf=="pct" else cv, fmtc, bg=bg)
        ws.cell(row=i,column=3).fill=PatternFill("solid",fgColor=bg); ws.cell(row=i,column=3).border=_b()
        _h(ws.cell(row=i,column=4),pl,bg=bg,fg=RX,bold=False,ha="left")
        _d(ws.cell(row=i,column=5), pv/100 if pf=="pct" and isinstance(pv,(int,float)) else pv,
           "0.0000%", bg=bg, fg=VD, bold=True)


# ══════════════════════════════════════════════════════════════════════════════
def _aba_fluxo(wb, resultado):
    ws = wb.create_sheet("Fluxo Financeiro")
    ws.sheet_view.showGridLines = False

    r  = resultado["resumo"]
    rc = resultado["receita"]
    df = resultado["fluxo"]

    # Título
    ws.merge_cells("A1:AC1")
    ws["A1"] = (f"FLUXO FINANCEIRO  |  {rc['nome']}  |  "
                f"C=R${r['Capital (C)']:,.0f}  T={r['Prazo (T)']}x  "
                f"i={r['Taxa ativa % a.m.']}%a.m.  CDI={r['CDI % a.m.']}%a.m.  "
                f"FPR={r['FPR']}  λ={r['λ Margem % a.m.']}%a.m.")
    ws["A1"].font  = Font(name="Calibri", bold=True, size=10, color=BR)
    ws["A1"].fill  = PatternFill("solid", fgColor=AZ)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 20

    # Grupos de colunas
    grupos = [
        ("t",                             1,  1, AZ),
        ("ATIVO — JUROS (z036/z043/F1023)",2,  5, AZ2),
        ("ATIVO — CAP./PMTA/SDA",          6, 10, AZ2),
        ("PASSIVO (z103/z109/z407/z161/z173)",11,18,"1A5C3A"),
        ("SPREAD/TARIFA (z175)",          19, 20, RX),
        ("TRIBUT. SP (z306/z311/ISS)",    21, 23, AM),
        ("PROVISÃO IPP/PE",               24, 27, "7B0000"),
        ("MARGENS (z182/z196)",           28, 29, RX),
        ("CUSTOS (F1044)",                30, 30, AM),
        ("BFB/IR-CS/FLA (F1048/z355/F1060/F1062)",31,35,"333333"),
    ]
    for lbl, c1, c2, cor in grupos:
        ws.merge_cells(f"{get_column_letter(c1)}2:{get_column_letter(c2)}2")
        _h(ws.cell(row=2, column=c1), lbl, bg=cor, sz=8)
        for cc in range(c1+1, c2+1):
            ws.cell(row=2,column=cc).fill = PatternFill("solid",fgColor=cor)
            ws.cell(row=2,column=cc).border = _b()
    ws.row_dimensions[2].height = 14

    hdrs = ["t","JA","EJA","SJA","EBA","LCA","ECA","SCA","PMTA","SDA",
            "EBP","EEBP","SEBP","LCP","ECP","SCP","PMTP","SDP",
            "SP","Tarifa",
            "BFA","PASEP+COF","ISS",
            "DPE","IPP Fase1","EPE","SIPP",
            "MG z182","MC z196",
            "f1044 Custos",
            "BFB","IR/CS","FLA","FLA VP","FLE"]
    for c, h in enumerate(hdrs, 1):
        _h(ws.cell(row=3, column=c), h, sz=8)
    ws.row_dimensions[3].height = 28

    df_cols = ["t","JA","EJA","SJA","EBA","LCA","ECA","SCA","PMTA","SDA",
               "EBP","EEBP","SEBP","LCP","ECP","SCP","PMTP","SDP",
               "SP","tarifa","BFA","pasep_cof","iss",
               "DPE","IPP1","EPE","SIPP",
               "MG","MC","f1044","BFB","ir_cs","FLA","FLA_VP","FLE"]

    T_ = r["Prazo (T)"]
    for idx, row in df.iterrows():
        r_ex = idx + 4
        bg = CZ if idx%2==0 else BR
        is_pe = (abs(row.get("PE", 0)) > 0)
        if is_pe: bg = AM_CL
        for c, col in enumerate(df_cols, 1):
            cel = ws.cell(row=r_ex, column=c)
            val = row.get(col, 0)
            if c == 1:
                _d(cel, int(val), "General", bg, ha="center")
            else:
                _d(cel, round(float(val), 4), "#,##0.0000", bg)

    # Linha de totais
    tr = len(df) + 4
    soma_c = {"JA","EBP","SP","pasep_cof","iss","DPE","IPP1","EPE",
              "MG","MC","f1044","BFB","ir_cs","FLA","FLA_VP","PMTA","PMTP","ECA","ECP"}
    for c, col in enumerate(df_cols, 1):
        cel = ws.cell(row=tr, column=c)
        if col in soma_c:
            _h(cel, round(float(df[col].sum()), 2), bg=AZ, sz=8)
            cel.number_format = "#,##0.00"
        else:
            _h(cel, "", bg=AZ, sz=8)

    lgs = [5]+[11]*9+[10]*8+[10]*4+[12]*4+[12]*2+[12]+[12]*2+[12]*3
    for i, w in enumerate(lgs[:len(df_cols)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ══════════════════════════════════════════════════════════════════════════════
def _aba_indicadores(wb, resultado):
    ws = wb.create_sheet("Indicadores")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 36

    r = resultado["resumo"]

    ws.merge_cells("A1:C1")
    ws["A1"] = "INDICADORES DE PERFORMANCE"
    ws["A1"].font = Font(name="Calibri", bold=True, size=12, color=BR)
    ws["A1"].fill = PatternFill("solid", fgColor=AZ)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 22

    GRUPOS = [
        ("SPREAD E MARGENS", [
            ("Total SP (spread bruto)", r["Total SP"], "R$", "Σ(PMTA−PMTP) no prazo"),
            ("Total MG (margem ganho)", r["Total MG"], "R$", "Σ(MG z182) no prazo"),
            ("SPVP", r["SPVP"], "R$", "SP trazido a VP pela curva SELIC"),
            ("MGVP (λ_mg)", r["MGVP (λ_mg)"], "R$", "MG VP — base dos indicadores de retorno"),
        ], AZ2),
        ("PROVISÃO E RISCO", [
            ("Total IPP Fase1", r["Total IPP Fase1"], "R$", "Provisão corrente mensal acumulada"),
            ("Total EPE (perda esperada)", r["Total EPE"], "R$", "Exigibilidade PE acumulada"),
            ("EPEVP", r["EPEVP"], "R$", "EPE a VP"),
            ("IR (índice de risco)", r["IR (risco)"], None, "EPEVP / SPVP — >1 = risco > spread"),
        ], "7B0000"),
        ("PERFORMANCE — RETORNO SOBRE CAPITAL", [
            ("MG % a.m.", r["MG % a.m."], "%", "λ mensal de margem"),
            ("MG % a.a.", r["MG % a.a."], "%", "λ anual = (1+λ_am)^12−1"),
            ("RSPLE % (F1064)", r["RSPLE %"], "%", "λ_aa / (FPR×K×(1+(F−1)×FCC)) × 100"),
            ("RAR_G % gestão (F1183)", r["RAR_G % (gestão, F1183)"], "%", "λ_aa / (FPR×Kp×(1+(F−1)×FCC)) × 100"),
            ("RAR % base (F1065)", r["RAR % (base)"], "%", "MGVP / CE × 100"),
            ("IE (eficiência)", r["IE (eficiência)"], None, "CFVP / SPVP — menor = mais eficiente"),
            ("IRE (rentabilidade)", r["IRE (rentabilidade)"], None, "MGVP / SPVP"),
            ("Duration (meses)", r["Duration (meses)"], None, "Duration ponderada por FLA VP"),
        ], RX),
        ("TAXAS IMPLÍCITAS — MÉTODOS ITERATIVOS", [
            ("TIR % a.a. (F1095)", r.get("TIR % a.a."), "%", "Taxa interna de retorno do Fluxo C"),
            ("Resultado PPS % a.m. (F1096)", r.get("Resultado PPS % a.m."), "%", "TPP — taxa do passivo que zera FLA VP"),
            ("FLA VP (Σ) → 0", r["FLA VP (Σ) → 0"], None, "Deve convergir para ~0 na taxa de equilíbrio"),
        ], VD),
    ]

    row = 2
    for grupo_nome, itens, cor in GRUPOS:
        ws.merge_cells(f"A{row}:C{row}")
        c = ws.cell(row=row, column=1, value=grupo_nome)
        c.font = Font(name="Calibri", bold=True, size=10, color=BR)
        c.fill = PatternFill("solid", fgColor=cor)
        c.alignment = Alignment(horizontal="left")
        c.border = _b()
        for cc in ["B","C"]: ws[f"{cc}{row}"].fill=PatternFill("solid",fgColor=cor); ws[f"{cc}{row}"].border=_b()
        ws.row_dimensions[row].height = 18
        row += 1

        for label, val, fmt, desc in itens:
            bg = CZ if row%2==0 else BR
            cel_a = ws.cell(row=row, column=1, value=label)
            cel_a.font = Font(name="Calibri", size=10, bold=True)
            cel_a.fill = PatternFill("solid", fgColor=bg)
            cel_a.border = _b()
            cel_a.alignment = Alignment(horizontal="left")

            if val is None: val = 0.0
            col_v = VD if (isinstance(val, float) and val > 0) else ("000000" if isinstance(val,float) and val==0 else "000000")
            cel_b = ws.cell(row=row, column=2)
            if fmt == "R$":
                _d(cel_b, val, "R$ #,##0.00", bg=bg, fg=col_v, bold=True)
            elif fmt == "%":
                _d(cel_b, val/100, "0.0000%", bg=bg, fg=col_v, bold=True)
            else:
                _d(cel_b, round(val,4) if isinstance(val,float) else val, None, bg=bg, bold=True)

            cel_c = ws.cell(row=row, column=3, value=desc)
            cel_c.font = Font(name="Calibri", size=9, italic=True, color="666666")
            cel_c.fill = PatternFill("solid", fgColor=bg)
            cel_c.border = _b()
            cel_c.alignment = Alignment(horizontal="left", wrap_text=True)
            ws.row_dimensions[row].height = 18
            row += 1
        row += 1
