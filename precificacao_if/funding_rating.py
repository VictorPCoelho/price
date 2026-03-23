"""
funding_rating.py
═════════════════
Custo de funding específico por rating.
Permite configurar manualmente ou carregar de arquivo.
"""
import io
import pandas as pd
from rating_model import RATINGS_ORDEM, PD_DEFAULTS_AA

# Defaults: custo de funding cresce com o risco do rating
# (quanto pior o rating, mais caro o funding para cobrir o risco)
FUNDING_DEFAULTS_AA = {
    # Letra A — funding prefixado base
    "A01": 6.50, "A02": 6.50, "A03": 6.75, "A04": 6.75, "A05": 7.00,
    # Letra B
    "B01": 7.00, "B02": 7.25, "B03": 7.25, "B04": 7.50, "B05": 7.50,
    # Letra C
    "C01": 7.75, "C02": 8.00, "C03": 8.00, "C04": 8.25, "C05": 8.50,
    # Letra D
    "D01": 8.50, "D02": 9.00, "D03": 9.00, "D04": 9.50, "D05":10.00,
    # Letra E
    "E01":10.00, "E02":10.50, "E03":11.00, "E04":11.50, "E05":12.00,
    # Letra F
    "F01":12.00, "F02":13.00, "F03":14.00, "F04":15.00, "F05":16.00,
    # Letra G
    "G":  18.00,
}

def taxa_am_de_aa(taxa_aa: float) -> float:
    return ((1 + taxa_aa/100)**(1/12) - 1) * 100

def funding_am_por_rating(funding_aa: dict) -> dict:
    """Converte dict {rating: taxa_aa} → {rating: taxa_am}."""
    return {cod: taxa_am_de_aa(v) for cod, v in funding_aa.items()}

def resumo_por_letra(funding_aa: dict) -> list:
    """Sumariza o custo médio por letra para exibição."""
    letras = list(dict.fromkeys(r[0] for r in RATINGS_ORDEM))
    rows = []
    for letra in letras:
        cods = [r for r in RATINGS_ORDEM if r.startswith(letra)]
        vals = [funding_aa.get(c, FUNDING_DEFAULTS_AA.get(c, 8.0)) for c in cods]
        rows.append({
            "Letra":    letra,
            "Ratings":  f"{cods[0]}–{cods[-1]}" if len(cods)>1 else cods[0],
            "Custo mín % a.a.": round(min(vals), 2),
            "Custo máx % a.a.": round(max(vals), 2),
            "Custo médio % a.a.": round(sum(vals)/len(vals), 2),
            "Custo médio % a.m.": round(taxa_am_de_aa(sum(vals)/len(vals)), 4),
        })
    return rows

def gerar_template_funding_csv() -> bytes:
    """Gera template CSV para upload de custo de funding por rating."""
    rows = []
    for cod in RATINGS_ORDEM:
        rows.append({
            "Rating": cod,
            "Custo_funding_pct_aa": FUNDING_DEFAULTS_AA.get(cod, 8.0),
            "Observacao": "Custo de captação % a.a. para este rating",
        })
    df = pd.DataFrame(rows)
    return df.to_csv(index=False, sep=";", decimal=",").encode("utf-8")

def gerar_template_funding_xlsx() -> bytes:
    """Gera template Excel para upload de custo de funding por rating."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Funding por Rating"
    CORES = {"A":"E8F5ED","B":"EEF2FA","C":"F3EFFE","D":"FFF3E0",
             "E":"FCEBEB","F":"FFD9D9","G":"FFCCCC"}
    hdrs = ["Rating","Letra","Custo funding % a.a.","Custo funding % a.m.","Observação"]
    for i, h in enumerate(hdrs, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A5C3A")
        c.alignment = Alignment(horizontal="center")
    for i, cod in enumerate(RATINGS_ORDEM, 2):
        letra = cod[0]
        custo_aa = FUNDING_DEFAULTS_AA.get(cod, 8.0)
        custo_am = taxa_am_de_aa(custo_aa)
        vals = [cod, letra, round(custo_aa, 2), round(custo_am, 6),
                "Edite 'Custo funding % a.a.'"]
        for j, val in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.fill = PatternFill("solid", fgColor=CORES.get(letra, "F5F5F5"))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

def carregar_funding_arquivo(conteudo: bytes, nome: str) -> tuple:
    """
    Carrega custo de funding por rating de xlsx ou csv.
    Retorna (dict {rating: custo_aa}, avisos[])
    """
    avisos = []
    try:
        if nome.lower().endswith((".xlsx", ".xls")):
            df = pd.read_excel(io.BytesIO(conteudo))
        else:
            for sep in [";", ",", "\t"]:
                try:
                    df = pd.read_csv(io.BytesIO(conteudo), sep=sep, decimal=",")
                    if len(df.columns) > 1: break
                except: pass

        # Normaliza nomes das colunas
        df.columns = [str(c).strip().lower().replace(" ","_") for c in df.columns]

        # Detecta coluna de rating
        col_rat = next((c for c in df.columns if "rating" in c), None)
        if not col_rat:
            raise ValueError("Coluna 'Rating' não encontrada.")

        # Detecta coluna de custo
        col_custo = next((c for c in df.columns
                          if any(x in c for x in ["custo","funding","taxa","cost"])), None)
        if not col_custo:
            raise ValueError("Coluna de custo não encontrada. Nomeie como 'Custo_funding_pct_aa'.")

        resultado = {}
        for _, row in df.iterrows():
            rat = str(row[col_rat]).strip().upper()
            if rat not in RATINGS_ORDEM: continue
            try:
                val = float(str(row[col_custo]).replace(",", "."))
                resultado[rat] = round(val, 4)
            except:
                avisos.append(f"Rating {rat}: valor inválido — usando default")

        # Ratings ausentes
        ausentes = [r for r in RATINGS_ORDEM if r not in resultado]
        if ausentes:
            avisos.append(f"{len(ausentes)} ratings sem custo no arquivo — usando defaults")
            for r in ausentes:
                resultado[r] = FUNDING_DEFAULTS_AA.get(r, 8.0)

        return resultado, avisos
    except Exception as e:
        raise ValueError(f"Erro ao ler arquivo: {e}")
